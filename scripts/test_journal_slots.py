import unittest,io,tempfile,zipfile,re,xml.etree.ElementTree as E
from pathlib import Path
from copy import copy,deepcopy
from openpyxl.utils import column_index_from_string
import openpyxl
from journal_slots import ensure_lesson_capacity,slot_columns,counter_value
from journal_input_controls import assert_template_controls,lesson_columns
import export_by_grade_subject as generator
import extract_journal_to_json as extract

class LessonCapacityTests(unittest.TestCase):
    def workbook(self,kind='MAIN'):
        w=openpyxl.load_workbook(Path(__file__).with_name(f'TEMPLATE_{kind}.xlsx'));self.addCleanup(w.close)
        s=w.active
        for merged in list(s.merged_cells.ranges):
            if merged.min_col>80:s.unmerge_cells(str(merged))
        s.delete_cols(81,max(s.max_column-80,1))
        for key in list(s.column_dimensions):
            if column_index_from_string(key)>80:del s.column_dimensions[key]
        for d in s.data_validations.dataValidation:
            d.sqref=' '.join(str(r) for r in d.sqref.ranges if r.max_col<=80)
        s.data_validations.dataValidation=[d for d in s.data_validations.dataValidation if d.sqref]
        return w
    def test_append_main_preserves_records_formulas_and_styles(self):
        w=self.workbook();s=w.active;source=w.copy_worksheet(s);source.data_validations=deepcopy(s.data_validations)
        s['D7']='既存本文';s['E9']='宿題';s['E11']='記録';s['D17']='https://example.com/video'
        s['FG2']=37;s['FG3']=2
        addresses=['D7','E9','E11','D17','B2','F2','B11']
        before={a:(s[a].value,copy(s[a]._style)) for a in addresses}
        ensure_lesson_capacity(s,source)
        self.assertEqual(len(lesson_columns(s)),21)
        self.assertEqual(before,{a:(s[a].value,s[a]._style) for a in addresses})
        self.assertEqual((s['HD2'].value,s['HD3'].value),(37,2))
        self.assertIn('$HD$2',s['GX2'].value)
        for col in range(82,203,10):
            self.assertIsNone(s.cell(7,col+2).value)
            self.assertIsNone(s.cell(9,col+3).value)
            self.assertIsNone(s.cell(11,col+3).value)
            self.assertIsNone(s.cell(17,col+2).value)
        assert_template_controls(s)
    def test_21st_controls_survive_save_reload_and_second_run(self):
        for kind in ['MAIN','X']:
            with self.subTest(kind=kind):
                w=self.workbook(kind);s=w.active;ensure_lesson_capacity(s)
                count=len(s.data_validations.dataValidation)
                self.assertEqual(ensure_lesson_capacity(s),0)
                self.assertEqual(len(s.data_validations.dataValidation),count)
                data=io.BytesIO();w.save(data);data.seek(0);loaded=openpyxl.load_workbook(data);self.addCleanup(loaded.close)
                final=loaded.active
                for addr in ['GT11','GT12','GZ20']:
                    ds=[d for d in final.data_validations.dataValidation if addr in d.sqref]
                    self.assertEqual(len(ds),1,addr);self.assertFalse(ds[0].showDropDown)
                self.assertEqual(len(lesson_columns(final)),21)
                self.assertEqual(generator.count_slots_in_template(final),21)
    def test_conflicting_content_in_new_area_stops(self):
        w=self.workbook();s=w.active;s['CF6']='既存のメモ'
        with self.assertRaisesRegex(ValueError,'already has content'):ensure_lesson_capacity(s)
    def test_new_month_21_slots_header_helper_and_existing_month_untouched(self):
        w=self.workbook();s=w.active;s.title='__TEMPLATE_MAIN__';s.sheet_state='hidden'
        old=w.create_sheet('2026-09');old['D7']='保持';old['F2']=37;old['B11']=10
        new=generator.create_month_sheet(w,s,2026,10)
        generator.set_header_cells(new,'本校','中３','英語',10,wb=w,year=2026)
        self.assertEqual(len(lesson_columns(new)),21)
        self.assertEqual(new['HD2'].value,38);self.assertEqual(old['D7'].value,'保持')
        self.assertIsNone(generator.create_month_sheet(w,s,2026,10))
    def test_extract_finds_18th_and_21st_and_previous_entry(self):
        w=self.workbook();s=w.active;ensure_lesson_capacity(s);s.title='2026-09'
        # Mimic Excel data_only loading a workbook whose formulas have no caches.
        for row in s:
            for c in row:
                if c.data_type=='f':c.value=None
        s['F2']=37;s['E3']=9;s['G3']=1
        for i in range(21):
            s.cell(11,2+10*i,i+1);s.cell(3,5+10*i,9)
        s.cell(7,174,'18列目の本文');s.cell(10,172,9)
        self.assertEqual(extract._find_block_col(s,6,21),202)
        self.assertEqual(extract._find_last_slot_col(s),202)
        self.assertEqual(extract.read_slot_header(s,202)['sessionNumber'],'57')
        prev=extract._read_prev_entry(w,s,2026,9,21,'S',6)
        self.assertIsNotNone(prev)
    def test_special_first_counter_uses_hd_without_corrupting_17th_header(self):
        w=self.workbook();s=w.active;ensure_lesson_capacity(s)
        s['F2']='特';s['HD2']=37;s['E3']=9;s['G3']=1;s['B11']=1;s['L11']=2
        s['P2']=None;s['O3']=None;s['Q3']=None
        self.assertEqual(extract.read_slot_header(s,12)['sessionNumber'],'37')
        self.assertEqual(counter_value(s,2),37)
        self.assertEqual(s['FF2'].value,'=$B$2')
    def test_overflow_is_rejected_before_dates_are_written(self):
        w=self.workbook('X');s=w.active;ensure_lesson_capacity(s)
        with self.assertRaisesRegex(ValueError,'授業枠不足'):generator.fill_sheet_x(s,9,[None]*22)
        self.assertIsNone(s['B11'].value)
    def test_main_and_x_schedule_write_through_slot_21(self):
        for kind,klass in [('MAIN','S'),('X','X')]:
            w=self.workbook(kind);s=w.active;ensure_lesson_capacity(s)
            events=[generator.Event(9,i+1,'月','18:00','1','3',klass,'英','3'+klass+'英','講師',False,False,i,1) for i in range(21)]
            if kind=='MAIN':
                generator.fill_sheet_main(s,9,{'S':events,'A':[],'B':[]})
                with self.assertRaisesRegex(ValueError,'授業枠不足'):generator.fill_sheet_main(s,9,{'S':events+[events[-1]],'A':[],'B':[]})
            else:generator.fill_sheet_x(s,9,events)
            self.assertEqual(s['GT11'].value,'21');self.assertEqual(s['GT14'].value,'講師')
    def test_preserving_patch_keeps_self_closing_rows_in_place(self):
        from journal_slot_patch import extend_file,NS
        w=self.workbook();s=w.active;s.row_dimensions[1].height=10;s['D7']='元の記録'
        with tempfile.TemporaryDirectory() as directory:
            src=Path(directory)/'source.xlsx';out=Path(directory)/'fixed.xlsx';w.save(src)
            with zipfile.ZipFile(src) as z:parts={n:z.read(n) for n in z.namelist()}
            part='xl/worksheets/sheet1.xml';raw=parts[part].decode()
            raw=re.sub(r'(<row\b[^>]*r="1"[^>]*)></row>',r'\1/>',raw)
            self.assertRegex(raw,r'<row\b[^>]*r="1"[^>]*/>')
            parts[part]=raw.encode()
            with zipfile.ZipFile(src,'w',zipfile.ZIP_DEFLATED) as z:
                for n,b in parts.items():z.writestr(n,b)
            extend_file(src,out,canonical=True)
            with zipfile.ZipFile(out) as z:final=E.fromstring(z.read(part))
            n={'s':NS}
            self.assertEqual(len(final.find('s:sheetData/s:row[@r="1"]',n)),0)
            for row in final.findall('s:sheetData/s:row',n):
                for c in row:self.assertTrue(c.get('r').endswith(row.get('r')))
            saved=openpyxl.load_workbook(out);self.addCleanup(saved.close)
            self.assertEqual(saved.active['D7'].value,'元の記録')

if __name__=='__main__':unittest.main()
