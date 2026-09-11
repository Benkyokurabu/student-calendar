"""Month generation failures must reach the workflow's publication boundary."""
import contextlib
import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import ci_extract_and_push as pipeline
import export_by_grade_subject as generator


class MonthFailureTests(unittest.TestCase):
    def exercise(self, count, kind='MAIN', existing=False):
        scripts = Path(__file__).parent.resolve()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            schedule = root / 'schedule.xlsx'
            book = openpyxl.Workbook()
            book.save(schedule)
            book.close()
            name = '本校中３' + ('X' if kind == 'X' else '') + '英語_2026.xlsx'
            annual = root / name
            book = openpyxl.load_workbook(scripts / f'TEMPLATE_{kind}.xlsx')
            book.active.title = f'__TEMPLATE_{kind}__'
            book.active.sheet_state = 'hidden'
            old = book.create_sheet('2026-09')
            old['D7'] = 'existing record'
            if existing:
                book.create_sheet('2026-10')['D7'] = 'already entered'
            book.save(annual)
            book.close()
            before = annual.read_bytes()
            klass = 'X' if kind == 'X' else 'S'
            events = [generator.Event(10, i + 1, '月', '18:00', '1', '3', klass,
                      '英', '3' + klass + '英', '講師', False, False, i, 1) for i in range(count)]
            original_cwd = os.getcwd()
            with patch.object(generator, 'pick_schedule_in_same_folder', return_value=(2026, 10, schedule)), \
                 patch.object(generator, 'choose_target_sheets', return_value=[('本校', 'sheet', object())]), \
                 patch.object(generator, 'collect_events', return_value=events), \
                 contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                if count > 21:
                    with self.assertRaisesRegex(RuntimeError, '授業枠不足') as raised:
                        pipeline.create_month_sheets(scripts, ['2026-10'], root)
                    self.assertIsInstance(raised.exception.__cause__, ValueError)
                    self.assertEqual(annual.read_bytes(), before)
                else:
                    result = pipeline.create_month_sheets(scripts, ['2026-10'], root)
                    self.assertEqual(result, 0 if existing else 1)
                    if existing:
                        self.assertEqual(annual.read_bytes(), before)
                    else:
                        saved = openpyxl.load_workbook(annual)
                        try:
                            self.assertEqual(saved['2026-09']['D7'].value, 'existing record')
                            self.assertEqual(saved['2026-10']['GT11'].value, '21')
                        finally:
                            saved.close()
            self.assertEqual(os.getcwd(), original_cwd)

    def test_main_22_stops_and_keeps_original(self):
        self.exercise(22)

    def test_x_22_stops_and_keeps_original(self):
        self.exercise(22, 'X')

    def test_main_21_is_still_generated(self):
        self.exercise(21)

    def test_x_21_is_still_generated(self):
        self.exercise(21, 'X')

    def test_existing_month_remains_a_successful_noop(self):
        self.exercise(21, existing=True)

    def test_existing_month_also_stops_when_schedule_exceeds_capacity(self):
        self.exercise(22, existing=True)

    def test_parallel_classes_are_not_counted_as_63_columns(self):
        events = [generator.Event(10, i + 1, '月', '18:00', '1', '3', 'S',
                  '英', '3S英', '講師', False, False, i, 1) for i in range(21)]
        pipeline.require_month_capacity(classes={k: events for k in ['S', 'A', 'B']})

    def test_pipeline_does_not_extract_after_month_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            scripts = Path(directory) / 'scripts'
            scripts.mkdir()
            (scripts / '_cloud_journal').mkdir()
            with patch.object(pipeline, '__file__', str(scripts / 'ci_extract_and_push.py')), \
                 patch.object(pipeline.sys, 'argv', ['ci_extract_and_push.py', '2026-10']), \
                 patch.object(pipeline, 'repair_hidden_template_validations'), \
                 patch.object(pipeline, 'create_month_sheets', side_effect=RuntimeError('授業枠不足')), \
                 patch.object(pipeline, 'run') as run, contextlib.redirect_stdout(io.StringIO()):
                with self.assertRaisesRegex(RuntimeError, '授業枠不足'):
                    pipeline.main()
                run.assert_not_called()


if __name__ == '__main__':
    unittest.main()
