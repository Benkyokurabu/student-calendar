"""Counters must follow the workbook, including manual operational adjustments."""
import unittest

import openpyxl
import extract_journal_to_json as journal


class WorkbookCounterTests(unittest.TestCase):
    def setUp(self):
        self.book = openpyxl.Workbook()
        self.book.active.title = "2026-08"
        self.book.active["F2"] = 1
        self.book.active["E3"] = 8
        self.book.active["B31"] = 24
        self.sheet = self.book.create_sheet("2026-09")
        self.sheet["F2"] = 28
        self.sheet["E3"] = 9
        self.sheet["G3"] = 2
        for col, day in [(2, 7), (12, 14), (22, 21)]:
            self.sheet.cell(31, col, day)

    def tearDown(self):
        self.book.close()

    def header(self, col):
        return journal.read_slot_header(self.sheet, col, wb=self.book, year=2026, month=9)

    def test_manual_start_is_not_recounted_from_previous_month(self):
        self.assertEqual(self.header(2), {"sessionNumber": "28", "monthNum": "9", "weekNum": "2"})

    def test_missing_formula_cache_uses_sheet_start_and_monthly_start(self):
        self.assertEqual(self.header(12), {"sessionNumber": "29", "monthNum": "9", "weekNum": "3"})

    def test_cached_or_explicit_target_values_are_preserved(self):
        self.sheet["P2"] = 35
        self.sheet["O3"] = 10
        self.sheet["Q3"] = 4
        self.assertEqual(self.header(12), {"sessionNumber": "35", "monthNum": "10", "weekNum": "4"})

    def test_partial_cache_only_fills_missing_fields(self):
        self.sheet["P2"] = 35
        self.assertEqual(self.header(12), {"sessionNumber": "35", "monthNum": "9", "weekNum": "3"})

    def test_special_slot_does_not_increment_regular_count(self):
        self.sheet["P2"] = "特"
        self.assertEqual(self.header(12)["sessionNumber"], "特")
        self.assertEqual(self.header(22)["sessionNumber"], "29")

    def test_special_first_slot_uses_backup_start(self):
        self.sheet["F2"] = "特"
        self.sheet.cell(2, 163, 28)
        self.assertEqual(self.header(12)["sessionNumber"], "28")

    def test_missing_start_does_not_invent_count_from_history(self):
        self.sheet["F2"] = None
        self.assertEqual(self.header(12)["sessionNumber"], "")

    def test_previous_entry_uses_the_same_workbook_header_rules(self):
        self.sheet["D27"] = "previous lesson"
        self.sheet["B30"] = 9
        previous = journal._read_prev_entry(self.book, self.sheet, 2026, 9, 2, "A", 26)
        self.assertEqual(previous["sessionNumber"], "28")
        self.assertEqual(previous["weekNum"], "2")


if __name__ == "__main__":
    unittest.main()
