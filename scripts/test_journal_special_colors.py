"""Special-only columns retain the exact same shading as unused lesson forms."""
from copy import copy
import io
from pathlib import Path
import unittest

import openpyxl
from openpyxl.styles import PatternFill

import export_by_grade_subject as generator
from journal_input_controls import is_unused_gray
from journal_slots import ensure_lesson_capacity


class SpecialLessonColorTests(unittest.TestCase):
    def sheet(self, kind="MAIN"):
        book = openpyxl.load_workbook(Path(__file__).with_name(f"TEMPLATE_{kind}.xlsx"))
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = "2026-12"
        ensure_lesson_capacity(sheet)
        generator.set_header_cells(sheet, "南教室", "中２", "数学", 12, wb=book, year=2026)
        return book, sheet

    def event(self, day, klass, special=False):
        return generator.Event(12, day, "月", "18:00", "1", "2", klass, "数",
                               "2" + klass + "数", "講師", special, False, day, 1)

    def generated(self, special_class="S"):
        book, sheet = self.sheet()
        classes = {klass: [self.event(7, klass), self.event(21, klass)] for klass in ("S", "A", "B")}
        classes[special_class].insert(1, self.event(14, special_class, True))
        generator.fill_sheet_main(sheet, 12, classes)
        return book, sheet

    def assert_special_column(self, sheet, special_class):
        for klass, top in [("S", 6), ("A", 26), ("B", 46)]:
            used = klass == special_class
            self.assertEqual(sheet.cell(top + 1, 12).value or "", "特" if used else "")
            self.assertEqual(str(sheet.cell(top + 5, 12).value or ""), "14" if used else "")
            # Compare the complete existing input-area fill mask, including
            # merged cells, with the same class's unused right-hand frame.
            for row in range(top + 1, top + 18):
                for offset in range(2, 9):
                    current = sheet.cell(row, 12 + offset).fill
                    unused = sheet.cell(row, 32 + offset).fill
                    if used:
                        self.assertFalse(is_unused_gray(current), (klass, row, offset))
                    else:
                        self.assertEqual(copy(current), copy(unused), (klass, row, offset))

    def test_special_s_a_b_blocks_match_unused_right_hand_frames_after_repair(self):
        for klass in ("S", "A", "B"):
            with self.subTest(special_class=klass):
                _, sheet = self.generated(klass)
                # Reproduce the old repair whitening an already-used column.
                for top in (6, 26, 46):
                    generator.clear_gray_block(sheet, top, 12)
                generator.restore_unused_slot_colors(sheet)
                self.assert_special_column(sheet, klass)

    def test_fills_survive_save_reload_and_repeated_repairs_without_changing_data(self):
        book, sheet = self.generated()
        sheet["D7"] = "通常授業の記録"
        sheet["N7"] = "特別授業の記録"
        sheet["X27"] = "通常Aクラスの記録"
        sheet["F8"].fill = PatternFill("solid", fgColor="FFFF0000")
        before = {c.coordinate: c.value for row in sheet for c in row if c.value not in (None, "")}
        for _ in range(2):
            generator.restore_unused_slot_colors(sheet)
        stream = io.BytesIO()
        book.save(stream)
        loaded = openpyxl.load_workbook(io.BytesIO(stream.getvalue()))
        self.addCleanup(loaded.close)
        result = loaded.active
        generator.restore_unused_slot_colors(result)
        self.assertEqual({c.coordinate: c.value for row in result for c in row if c.value not in (None, "")}, before)
        self.assertEqual(result["F8"].fill.fgColor.rgb, "FFFF0000")
        self.assert_special_column(result, "S")

    def test_regular_five_lessons_remain_writable_and_sixth_is_shaded(self):
        _, sheet = self.sheet()
        classes = {klass: [self.event(day, klass) for day in (1, 7, 14, 21, 28)] for klass in ("S", "A", "B")}
        generator.fill_sheet_main(sheet, 12, classes)
        generator.restore_unused_slot_colors(sheet)
        for top in (6, 26, 46):
            self.assertFalse(is_unused_gray(sheet.cell(top + 1, 44).fill))
            self.assertTrue(is_unused_gray(sheet.cell(top + 1, 54).fill))

    def test_existing_populated_class_is_not_shaded_by_an_inconsistent_special_label(self):
        _, sheet = self.generated()
        sheet["L27"] = "A"
        sheet["L31"] = 15
        sheet["N27"] = "既存の記録を保持"
        generator.restore_unused_slot_colors(sheet)
        self.assertFalse(is_unused_gray(sheet["N27"].fill))
        self.assertEqual(sheet["N27"].value, "既存の記録を保持")

    def test_x_lessons_keep_existing_colors_and_capacity(self):
        _, sheet = self.sheet("X")
        generator.fill_sheet_x(sheet, 12, [self.event(7, "X"), self.event(14, "X", True)])
        generator.restore_unused_slot_colors(sheet)
        self.assertFalse(is_unused_gray(sheet["D7"].fill))
        self.assertFalse(is_unused_gray(sheet["N7"].fill))
        self.assertTrue(is_unused_gray(sheet["X7"].fill))


if __name__ == "__main__":
    unittest.main()
