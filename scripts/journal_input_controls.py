"""Journal input controls: preserve custom rules, cover real lesson slots, reject broken templates."""
from copy import copy
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.datavalidation import DataValidation


def lesson_columns(ws):
    return [cell.column for row in ws.iter_rows(min_row=6,max_row=6) for cell in row if cell.value == 'クラス']


def is_unused_gray(fill):
    color = fill.fgColor
    return fill.patternType == 'solid' and color.type == 'rgb' and str(color.rgb)[-6:].upper() == 'D9D9D9'


def assert_template_controls(ws):
    required = ['B12','B11','B10','G3','F2','B2','E3','I2','I3','B7','H20','B14']
    for address in required:
        matches = [d for d in ws.data_validations.dataValidation if address in d.sqref]
        if len(matches) != 1 or matches[0].type != 'list' or not matches[0].formula1 or matches[0].showDropDown:
            raise ValueError(f'{ws.title}!{address}: 正規テンプレートのプルダウンが欠落・重複・非表示です')
    progress = next(d for d in ws.data_validations.dataValidation if 'H20' in d.sqref)
    if not all(symbol in progress.formula1 for symbol in ['＋','－','±']):
        raise ValueError(f'{ws.title}!H20: 進捗候補が不正です')


def ensure_controls(dst, source=None):
    """Use first-slot canonical rules, extend by actual headers, preserve existing custom lists."""
    source = source if source is not None else dst
    assert_template_controls(source)
    prototypes=[]
    for d in source.data_validations.dataValidation:
        first=[r for r in d.sqref.ranges if r.min_col <= 10]
        for r in first:
            # Single-cell rules from a repaired template also need to repeat.
            repeated = r.min_row >= 7 or (r.min_row == 2 and r.min_col == 6) or (r.min_row == 3 and r.min_col == 7)
            prototypes.append((copy(d),str(r),repeated))
    covered=set()
    for d in dst.data_validations.dataValidation:
        for r in d.sqref.ranges:
            covered.update((row,col) for row in range(r.min_row,r.max_row+1) for col in range(r.min_col,r.max_col+1))
    added=0
    for d, address, repeated in prototypes:
        new=copy(d);new.sqref='';new.showDropDown=False
        for left in lesson_columns(dst) if repeated else [2]:
            r=CellRange(address);r.shift(col_shift=left-2)
            for row in range(r.min_row,r.max_row+1):
                for col in range(r.min_col,r.max_col+1):
                    if (row,col) not in covered:
                        new.add(dst.cell(row,col).coordinate);covered.add((row,col));added+=1
        if new.sqref:dst.add_data_validation(new)
    assert_template_controls(dst)
    return added
