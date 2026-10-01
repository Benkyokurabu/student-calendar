import unittest
from unittest.mock import patch

from openpyxl import Workbook

import extract_journal_to_json as journal


class SameDayJournalColumnsTests(unittest.TestCase):
    def setUp(self):
        self.first = {
            "date": "2026-09-23", "time": "4:45～6:15", "grade": "j1",
            "class": "B", "subject": "math", "campus": "minami",
            "groupKey": "minami_j1_B_math", "room": "3", "special": True,
        }
        self.second = self.first | {"time": "6:35～8:05", "special": False}
        self.wb = Workbook()
        ws = self.wb.active
        ws.title = "2026-09"
        for col, content, page in (
            (22, "方程式、基礎計算練習", "定期テスト対策p14"),
            (32, "まとめ・中間対策", "定期テスト対策p15"),
        ):
            ws.cell(51, col, 23)
            ws.cell(47, col + 2, content)
            ws.cell(48, col + 4, page)
        ws.cell(2, 36, "特")

    def test_earlier_lesson_reads_left_column_and_later_reads_right(self):
        positions = journal.build_same_day_positions([self.second, self.first])
        self.assertEqual(positions[journal.make_entry_key(self.first)], (0, 2))
        self.assertEqual(positions[journal.make_entry_key(self.second)], (1, 2))

        class Cache:
            def get_workbook(_self, _filename):
                return self.wb

        with patch.object(journal, "read_slot_header", return_value={}), \
             patch.object(journal, "_read_prev_entry", return_value=None):
            first = journal.extract_entry_for_event(
                self.first, 4, Cache(),
                same_day_position=positions[journal.make_entry_key(self.first)],
            )
            second = journal.extract_entry_for_event(
                self.second, 3, Cache(),
                same_day_position=positions[journal.make_entry_key(self.second)],
            )

        self.assertEqual((first["content"], first["page"]),
                         ("方程式、基礎計算練習", "定期テスト対策p14"))
        self.assertEqual((second["content"], second["page"]),
                         ("まとめ・中間対策", "定期テスト対策p15"))

    def test_column_count_mismatch_stops_instead_of_guessing(self):
        self.wb.active.cell(51, 32).value = None

        class Cache:
            def get_workbook(_self, _filename):
                return self.wb

        with self.assertRaisesRegex(ValueError, "日付列数が一致しません"):
            journal.extract_entry_for_event(
                self.second, 3, Cache(), same_day_position=(1, 2)
            )

    def test_same_start_time_stops_instead_of_guessing(self):
        simultaneous = self.second | {"time": "4:45～7:00", "room": "4"}
        with self.assertRaisesRegex(ValueError, "開始時刻が重複"):
            journal.build_same_day_positions([self.first, simultaneous])


if __name__ == "__main__":
    unittest.main()
