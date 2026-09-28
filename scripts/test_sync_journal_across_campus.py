import unittest
from openpyxl import Workbook

from sync_journal_across_campus import (
    FIELD_OFFSETS, ONLINE_PAIR_SET, build_online_pairs, compare_block_fields,
    compare_x_block_fields, read_block, write_block,
)


class CompareBlockFieldsTests(unittest.TestCase):
    def test_all_current_x_classes_are_configured(self):
        self.assertIn(("e5", "X", "jp"), ONLINE_PAIR_SET)
        self.assertIn(("e6", "X", "jp"), ONLINE_PAIR_SET)

    def test_new_x_class_is_paired_when_both_campuses_share_a_lesson(self):
        events = [
            {"date": "2026-10-12", "time": "4:55～6:15", "grade": "e6",
             "class": "X", "subject": "arith", "campus": campus,
             "faceToFace": False}
            for campus in ("hon", "minami")
        ]
        self.assertEqual(1, len(build_online_pairs(events, "2026-10")))

    def test_x_class_copies_absence_progress_and_note(self):
        hon = Workbook().active
        minami = Workbook().active
        hon["C18"] = "欠席者"
        hon["H20"] = "+"
        hon["I20"] = "2"
        hon["C21"] = "次回確認"
        hon_block = read_block(hon, 6, 2)
        minami_block = read_block(minami, 6, 2)

        to_hon, to_minami, conflicts = compare_block_fields(
            hon_block, minami_block, FIELD_OFFSETS,
        )
        self.assertEqual({}, to_hon)
        self.assertEqual([], conflicts)
        self.assertEqual(
            {"absence1": "欠席者", "curriculumSign": "+",
             "curriculumValue": "2", "note": "次回確認"},
            to_minami,
        )
        write_block(minami, 6, 2, to_minami)
        self.assertEqual(hon_block, read_block(minami, 6, 2))

    def test_x_class_correction_replaces_old_value_on_other_campus(self):
        old = {"content": "旧内容"}
        to_hon, to_minami, conflicts = compare_x_block_fields(
            {"content": "訂正後"}, {"content": "旧内容"}, old, old,
        )
        self.assertEqual({}, to_hon)
        self.assertEqual({"content": "訂正後"}, to_minami)
        self.assertEqual([], conflicts)

    def test_x_class_clear_is_propagated(self):
        old = {"note": "旧備考"}
        to_hon, to_minami, conflicts = compare_x_block_fields(
            {"note": ""}, {"note": "旧備考"}, old, old,
        )
        self.assertEqual({}, to_hon)
        self.assertEqual({"note": ""}, to_minami)
        self.assertEqual([], conflicts)
        sheet = Workbook().active
        sheet["C21"] = "旧備考"
        write_block(sheet, 6, 2, to_minami)
        self.assertEqual("", sheet["C21"].value)

    def test_x_class_simultaneous_different_edits_remain_conflicts(self):
        old = {"content": "元の内容"}
        to_hon, to_minami, conflicts = compare_x_block_fields(
            {"content": "本校の訂正"}, {"content": "南教室の訂正"}, old, old,
        )
        self.assertEqual({}, to_hon)
        self.assertEqual({}, to_minami)
        self.assertEqual(["content"], conflicts)

    def test_legacy_pairs_keep_existing_field_scope(self):
        hon = {"content": "本文", "absence1": "欠席者"}
        to_hon, to_minami, conflicts = compare_block_fields(hon, {})
        self.assertEqual({}, to_hon)
        self.assertEqual({"content": "本文"}, to_minami)
        self.assertEqual([], conflicts)

    def test_partially_filled_destination_receives_only_missing_fields(self):
        hon = {
            "content": "現代文",
            "page": "12-15",
            "homework1": "漢字",
            "homework2": "",
            "report": "",
            "recordingUrl": "https://example.test/video",
        }
        minami = {
            "content": "",
            "page": "",
            "homework1": "",
            "homework2": "",
            "report": "",
            "recordingUrl": "https://example.test/video",
        }

        to_hon, to_minami, conflicts = compare_block_fields(hon, minami)

        self.assertEqual({}, to_hon)
        self.assertEqual(
            {"content": "現代文", "page": "12-15", "homework1": "漢字"},
            to_minami,
        )
        self.assertEqual([], conflicts)

    def test_complementary_fields_are_copied_in_both_directions(self):
        hon = {"content": "現代文", "page": "", "homework1": "", "homework2": "", "report": "", "recordingUrl": ""}
        minami = {"content": "", "page": "20", "homework1": "", "homework2": "", "report": "", "recordingUrl": ""}

        to_hon, to_minami, conflicts = compare_block_fields(hon, minami)

        self.assertEqual({"page": "20"}, to_hon)
        self.assertEqual({"content": "現代文"}, to_minami)
        self.assertEqual([], conflicts)

    def test_different_nonempty_values_are_never_overwritten(self):
        hon = {"content": "本校の記入", "page": "", "homework1": "", "homework2": "", "report": "", "recordingUrl": ""}
        minami = {"content": "南の記入", "page": "", "homework1": "", "homework2": "", "report": "", "recordingUrl": ""}

        to_hon, to_minami, conflicts = compare_block_fields(hon, minami)

        self.assertEqual({}, to_hon)
        self.assertEqual({}, to_minami)
        self.assertEqual(["content"], conflicts)


if __name__ == "__main__":
    unittest.main()
