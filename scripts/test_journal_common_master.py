"""Shared masters must update new months while retaining existing journal records."""
from copy import copy
import contextlib
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

import ci_extract_and_push as pipeline
import export_by_grade_subject as generator
import extract_journal_to_json as extract
from journal_slots import ensure_lesson_capacity, slot_columns, print_area_ranges

SCRIPTS = Path(__file__).resolve().parent


class CommonMasterTests(unittest.TestCase):
    def fixture(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        root = Path(folder.name)
        common = root / generator.COMMON_MAIN_TEMPLATE_NAME
        book = openpyxl.load_workbook(SCRIPTS / "TEMPLATE_MAIN.xlsx")
        self.addCleanup(book.close)
        source = book.active
        source.title = "共通マスター"
        ensure_lesson_capacity(source)
        for left in slot_columns(source):
            source.cell(2, left, "本校")
            source.cell(2, left + 7, "中３")
            source.cell(3, left + 7, "英語")
            source.cell(2, left + 4, 789)
        source["P2"] = "特"
        source["D6"] = "共通レイアウト"
        source.column_dimensions.group("B", "D")
        source.column_dimensions["B"].width = 9.25
        source.column_dimensions["B"].hidden = False
        source.row_dimensions[7].height = 32
        source["D7"].fill = PatternFill("solid", fgColor="FFF0E0")
        source["D8"].comment = Comment("マスターの注記", "User")
        source.print_area = "B2:HB65"
        source.print_title_rows = "1:5"
        source.page_setup.orientation = "landscape"
        source.page_margins.left = 0.4
        source.oddHeader.center.text = "授業日誌"
        source.conditional_formatting.add("D7:D8", CellIsRule(
            operator="equal", formula=['"要確認"'],
            fill=PatternFill("solid", fgColor="FFFF00")))
        book.save(common)
        return root, common, source

    def annual(self, kind="MAIN"):
        book = openpyxl.load_workbook(SCRIPTS / f"TEMPLATE_{kind}.xlsx")
        self.addCleanup(book.close)
        hidden = book.active
        hidden.title = f"__TEMPLATE_{kind}__"
        hidden.sheet_state = "hidden"
        hidden["D6"] = "従来のレイアウト"
        previous = book.create_sheet("2026-11")
        previous["D7"] = "保存済みの授業内容"
        previous["D7"].fill = PatternFill("solid", fgColor="FF112233")
        previous["F2"] = 40
        previous["E3"] = 11
        previous["G3"] = 1
        previous["B11"] = 2
        previous["L11"] = 9
        previous["V11"] = 30
        previous["Y3"] = 12  # The last lesson is accounted for in December.
        return book, hidden, previous

    def events(self):
        return [generator.Event(12, day, "月", "18:00", "1", "2", "S", "数",
                                "2S数", "講師", special, False, index, 1)
                for index, (day, special) in enumerate([(7, False), (14, True), (21, False)])]

    def test_every_frame_replaces_copied_labels_and_counts_for_all_targets(self):
        root, _, _ = self.fixture()
        for campus, grade, subject in [("南教室", "中２", "数学"), ("本校", "小５", "算数"),
                                        ("南教室", "中１", "国語"), ("本校", "中３", "理科"),
                                        ("南教室", "", "英語補講")]:
            with self.subTest(campus=campus, grade=grade, subject=subject):
                book, hidden, _ = self.annual()
                sheet = generator.create_month_sheet(book, hidden, 2026, 12,
                                                      template_path=root / "TEMPLATE_MAIN.xlsx")
                generator.set_header_cells(sheet, campus, grade, subject, 12, wb=book, year=2026)
                generator.fill_sheet_main(sheet, 12, {"S": self.events(), "A": [], "B": []})
                for left in slot_columns(sheet):
                    self.assertEqual(sheet.cell(2, left).value, campus)
                    self.assertEqual(sheet.cell(2, left + 7).value, grade)
                    self.assertEqual(sheet.cell(3, left + 7).value, subject)
                    self.assertIsNone(sheet.cell(3, left + 3).value)
                    self.assertIsNone(sheet.cell(3, left + 5).value)
                self.assertEqual(sheet["Z2"].data_type, "f")
                # Production extraction uses data_only=True; a freshly saved
                # formula has no cached result until Excel recalculates it.
                sheet["Z2"] = None
                self.assertEqual([extract.read_slot_header(sheet, left, wb=book, year=2026, month=12)
                                  ["sessionNumber"] for left in (2, 12, 22)], ["42", "特", "43"])

    def test_new_month_copies_common_layout_controls_and_printing_after_reload(self):
        root, _, source = self.fixture()
        book, hidden, previous = self.annual()
        previous_style = copy(previous["D7"]._style)
        previous_values = [(c.coordinate, c.value) for row in previous for c in row if c.value is not None]
        sheet = generator.create_month_sheet(book, hidden, 2026, 12,
                                              template_path=root / "TEMPLATE_MAIN.xlsx")
        self.assertEqual(sheet["D6"].value, "共通レイアウト")
        self.assertEqual(hidden["D6"].value, "従来のレイアウト")
        self.assertEqual(previous["D7"]._style, previous_style)
        self.assertEqual([(c.coordinate, c.value) for row in previous for c in row if c.value is not None], previous_values)
        output = root / "annual.xlsx"
        book.save(output)
        saved = openpyxl.load_workbook(output)
        self.addCleanup(saved.close)
        result = saved["2026-12"]
        self.assertEqual(saved["2026-11"]["D7"].value, "保存済みの授業内容")
        self.assertEqual(saved["2026-11"]["D7"]._style, previous_style)
        self.assertEqual(set(map(str, result.merged_cells.ranges)), set(map(str, source.merged_cells.ranges)))
        self.assertEqual({key: dict(value) for key, value in result.column_dimensions.items()},
                         {key: dict(value) for key, value in source.column_dimensions.items()})
        self.assertEqual(result.row_dimensions[7].height, 32)
        self.assertEqual(copy(result["D7"].fill), copy(source["D7"].fill))
        self.assertEqual(result["D8"].comment.text, source["D8"].comment.text)
        self.assertEqual(ET.tostring(result.data_validations.to_tree()), ET.tostring(source.data_validations.to_tree()))
        self.assertEqual(len(result.conditional_formatting), 1)
        self.assertEqual([str(area) for area in print_area_ranges(result)], ["B2:HB65"])
        self.assertEqual(result.print_title_rows, "$1:$5")
        self.assertEqual(result.page_setup.orientation, "landscape")
        self.assertEqual(result.page_margins.left, 0.4)
        self.assertEqual(result.oddHeader.center.text, "授業日誌")

    def test_unregistered_editing_copy_keeps_existing_main_layout(self):
        root, common, _ = self.fixture()
        common.rename(root / "通常用_新マスター_編集中.xlsx")
        book, hidden, _ = self.annual()
        sheet = generator.create_month_sheet(book, hidden, 2026, 12,
                                              template_path=root / "TEMPLATE_MAIN.xlsx")
        self.assertEqual(sheet["D6"].value, "従来のレイアウト")

    def test_x_layout_remains_separate_even_when_main_common_master_exists(self):
        root, _, _ = self.fixture()
        book, hidden, _ = self.annual("X")
        sheet = generator.create_month_sheet(book, hidden, 2026, 12,
                                              template_path=root / "TEMPLATE_X.xlsx")
        self.assertEqual(sheet["D6"].value, "従来のレイアウト")
        self.assertEqual(len(sheet.merged_cells.ranges), len(hidden.merged_cells.ranges))

    def test_invalid_common_master_stops_without_adding_a_month(self):
        root, common, _ = self.fixture()
        broken = openpyxl.load_workbook(common)
        self.addCleanup(broken.close)
        broken.create_sheet("授業記録が混じったシート")
        broken.save(common)
        book, hidden, previous = self.annual()
        before = book.sheetnames[:]
        with self.assertRaisesRegex(ValueError, "1シート"):
            generator.create_month_sheet(book, hidden, 2026, 12,
                                         template_path=root / "TEMPLATE_MAIN.xlsx")
        self.assertEqual(book.sheetnames, before)
        self.assertEqual(previous["D7"].value, "保存済みの授業内容")

    def test_existing_month_is_not_replaced_or_revalidated(self):
        root, common, _ = self.fixture()
        common.write_bytes(b"invalid master must not affect an existing month")
        book, hidden, previous = self.annual()
        self.assertIsNone(generator.create_month_sheet(book, hidden, 2026, 11,
                                                       template_path=root / "TEMPLATE_MAIN.xlsx"))
        self.assertEqual(previous["D7"].value, "保存済みの授業内容")

    def test_cloud_pipeline_uses_registered_common_master_for_existing_year_file(self):
        root, common, _ = self.fixture()
        scripts = root / "scripts"
        scripts.mkdir()
        (scripts / common.name).write_bytes(common.read_bytes())
        for name in ("TEMPLATE_MAIN.xlsx", "TEMPLATE_X.xlsx"):
            (scripts / name).write_bytes((SCRIPTS / name).read_bytes())
        journal = root / "journal"
        journal.mkdir()
        annual = journal / "南教室中２数学_2026.xlsx"
        book, _, _ = self.annual()
        book.save(annual)
        schedule = root / "schedule.xlsx"
        schedule_book = openpyxl.Workbook()
        schedule_book.save(schedule)
        schedule_book.close()
        with patch.object(generator, "pick_schedule_in_same_folder", return_value=(2026, 12, schedule)), \
             patch.object(generator, "choose_target_sheets", return_value=[("南教室", "教務", object())]), \
             patch.object(generator, "collect_events", return_value=self.events()), \
             contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(pipeline.create_month_sheets(scripts, ["2026-12"], journal), 1)
        saved = openpyxl.load_workbook(annual)
        self.addCleanup(saved.close)
        self.assertEqual(saved["2026-12"]["D6"].value, "共通レイアウト")
        self.assertEqual(saved["2026-12"]["HA3"].value, "数学")
        self.assertEqual(saved["2026-11"]["D7"].value, "保存済みの授業内容")


if __name__ == "__main__":
    unittest.main()
