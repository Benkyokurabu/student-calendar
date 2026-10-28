"""Month/week retirement must preserve lesson contents and annual/carryover counts."""
from copy import copy
import io
from pathlib import Path
import unittest

import openpyxl

import export_by_grade_subject as generator
import extract_journal_to_json as extract
from journal_slots import ensure_lesson_capacity, slot_columns
from journal_month_week import (
    clear_month_week_headers, slot_accounting_month, MONTH_WEEK_BLANK_FROM,
)


class BlankMonthWeekTests(unittest.TestCase):
    def workbook(self, kind="MAIN", title="2026-11"):
        book = openpyxl.load_workbook(Path(__file__).with_name(f"TEMPLATE_{kind}.xlsx"))
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = title
        ensure_lesson_capacity(sheet)
        return book, sheet

    def event(self, day, *, special=False, carry=False):
        return generator.Event(11, day, "月", "18:00", "1", "3", "S", "数",
                               "3S数", "講師", special, carry, 1, 1)

    def assert_blank_headers(self, sheet):
        for left in slot_columns(sheet):
            for offset in (3, 5):
                self.assertIsNone(sheet.cell(3, left + offset).value)

    def test_new_main_and_x_and_repair_keep_all_21_headers_blank(self):
        for kind in ("MAIN", "X"):
            with self.subTest(kind=kind):
                book, sheet = self.workbook(kind)
                generator.set_header_cells(sheet, "本校", "中３", "数学", 11, wb=book, year=2026)
                events = [self.event(2), self.event(9, special=True), self.event(16)]
                if kind == "MAIN":
                    generator.fill_sheet_main(sheet, 11, {"S": events, "A": [], "B": []})
                else:
                    generator.fill_sheet_x(sheet, 11, events)
                generator.patch_counter_formulas(sheet)
                self.assert_blank_headers(sheet)
                self.assertEqual(sheet["P2"].value, "特")
                self.assertEqual(sheet["V11"].value, "16")
                self.assertEqual(slot_accounting_month(sheet, 2), 11)

    def test_clear_preserves_records_styles_dates_annual_and_controls_after_reload(self):
        book, sheet = self.workbook()
        sheet["D7"] = "授業内容を保持"
        sheet["B11"] = "2"
        sheet["F2"] = 46
        sheet["E3"] = 11
        sheet["G3"] = 1
        style = copy(sheet["E3"]._style)
        controls = str(sheet.data_validations)
        clear_month_week_headers(sheet)
        buffer = io.BytesIO()
        book.save(buffer)
        saved = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
        self.addCleanup(saved.close)
        result = saved.active
        self.assert_blank_headers(result)
        self.assertEqual(result["D7"].value, "授業内容を保持")
        self.assertEqual(result["B11"].value, "2")
        self.assertEqual(result["F2"].value, 46)
        self.assertEqual(result["E3"]._style, style)
        self.assertEqual(str(result.data_validations), controls)
        self.assertTrue(result.column_dimensions["HD"].hidden)
        self.assertEqual(result["F3"].value, "月")
        self.assertEqual(result["H3"].value, "週")

    def test_future_existing_month_is_cleared_but_historical_month_is_unchanged(self):
        for title, clear in [("2026-09", False), ("2026-11", True)]:
            with self.subTest(title=title):
                book, sheet = self.workbook(title=title)
                sheet["E3"] = 9 if not clear else 11
                sheet["G3"] = 1
                clear_month_week_headers(sheet, from_date=MONTH_WEEK_BLANK_FROM)
                if clear:
                    self.assert_blank_headers(sheet)
                else:
                    self.assertEqual(sheet["E3"].value, 9)
                    self.assertEqual(sheet["G3"].value, 1)

    def test_october_boundary_preserves_old_lessons_and_clears_october_28_onwards(self):
        book, sheet = self.workbook(title="2026-10")
        sheet["E3"] = 10
        sheet["G3"] = 1
        sheet["B31"] = 23
        sheet["L31"] = 28
        sheet["O3"] = 10
        sheet["Q3"] = "=G3+1"
        clear_month_week_headers(sheet, from_date=MONTH_WEEK_BLANK_FROM)
        self.assertEqual(sheet["E3"].value, 10)
        self.assertEqual(sheet["G3"].value, 1)
        self.assertIsNone(sheet["O3"].value)
        self.assertIsNone(sheet["Q3"].value)
        generator.patch_counter_formulas(sheet)
        self.assertIsNone(sheet["Q3"].value)

    def test_legacy_carryover_month_survives_clear_and_repeated_updates(self):
        book, sheet = self.workbook(title="2026-11")
        sheet["F2"] = 46
        sheet["E3"] = 11
        sheet["G3"] = 1
        sheet["O3"] = 12
        sheet["B11"] = "23"
        sheet["L11"] = "30"
        for _ in range(2):
            clear_month_week_headers(sheet)
            generator.patch_counter_formulas(sheet)
        self.assertEqual(slot_accounting_month(sheet, 12), 12)
        self.assertEqual(generator.count_regular_slots_in_sheet(sheet, 11), 1)
        self.assertEqual(generator.get_last_session_number(sheet, 11), 46)
        self.assertEqual(generator.compute_annual_start(book, 2026, 12), 47)

    def test_new_carryover_retains_next_month_accounting(self):
        book, sheet = self.workbook()
        generator.set_header_cells(sheet, "南教室", "中３", "数学", 11, wb=book, year=2026)
        generator.fill_sheet_main(sheet, 11, {"S": [self.event(23), self.event(30, carry=True)], "A": [], "B": []})
        self.assert_blank_headers(sheet)
        self.assertEqual(slot_accounting_month(sheet, 12), 12)
        self.assertEqual(generator.count_regular_slots_in_sheet(sheet, 11), 1)

    def test_extraction_retains_annual_numbers_with_blank_month_week_and_missing_caches(self):
        book, sheet = self.workbook()
        sheet["F2"] = 46
        sheet["E3"] = 11
        sheet["G3"] = 1
        for left, day in [(2, 2), (12, 9), (22, 16)]:
            sheet.cell(11, left, day)
            sheet.cell(3, left + 3, 11)
        clear_month_week_headers(sheet)
        generator.patch_counter_formulas(sheet)
        sheet["P2"] = None
        sheet["Z2"] = None
        self.assertEqual(extract.read_slot_header(sheet, 22, wb=book, year=2026, month=11),
                         {"sessionNumber": "48", "monthNum": "11", "weekNum": "3"})

    def test_extraction_skips_special_and_carryover_with_blank_headers(self):
        book, sheet = self.workbook()
        generator.set_header_cells(sheet, "本校", "中３", "数学", 11, wb=book, year=2026)
        events = [self.event(2, special=True), self.event(9), self.event(30, carry=True)]
        generator.fill_sheet_main(sheet, 11, {"S": events, "A": [], "B": []})
        sheet["P2"] = None
        sheet["Z2"] = None
        for left, number, month, week in [(2, "特", "11", "特"), (12, "1", "11", "1"), (22, "2", "12", "1")]:
            self.assertEqual(extract.read_slot_header(sheet, left, wb=book, year=2026, month=11),
                             {"sessionNumber": number, "monthNum": month, "weekNum": week})

    def test_calendar_output_is_unchanged_after_clearing_manual_header_adjustments(self):
        book, sheet = self.workbook()
        sheet["F2"] = 46
        sheet["E3"] = 11
        sheet["G3"] = 2
        sheet["O3"] = 11
        sheet["Q3"] = 4
        sheet["Y3"] = 12
        sheet["AA3"] = 1
        for left, day in [(2, 16), (12, 23), (22, 30)]:
            sheet.cell(11, left, day)
        before = [extract.read_slot_header(sheet, left, wb=book, year=2026, month=11)
                  for left in (2, 12, 22)]
        clear_month_week_headers(sheet)
        after = [extract.read_slot_header(sheet, left, wb=book, year=2026, month=11)
                 for left in (2, 12, 22)]
        self.assertEqual(after, before)
        self.assert_blank_headers(sheet)


if __name__ == "__main__":
    unittest.main()
