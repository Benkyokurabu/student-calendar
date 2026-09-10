import unittest
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from journal_input_controls import ensure_controls, assert_template_controls, is_unused_gray
from export_by_grade_subject import clear_gray_block,create_month_sheet

class JournalInputControlsTests(unittest.TestCase):
 def reference(self):
  w=openpyxl.load_workbook(Path(__file__).with_name('TEMPLATE_X.xlsx'));self.addCleanup(w.close);return w.worksheets[0]
 def test_missing_template_rejected(self):
  with self.assertRaisesRegex(ValueError,'プルダウン'):assert_template_controls(openpyxl.Workbook().active)
 def test_partial_template_rejected(self):
  s=self.reference();s.data_validations.dataValidation=[d for d in s.data_validations.dataValidation if 'H20' not in d.sqref]
  with self.assertRaisesRegex(ValueError,'H20'):assert_template_controls(s)
 def test_alpha_variants_clear_without_touching_values_or_other_colors(self):
  for color in ['00D9D9D9','FFD9D9D9']:
   s=openpyxl.Workbook().active;s['D7']='本文';s['D7'].fill=PatternFill('solid',fgColor=color);s['F8'].fill=PatternFill('solid',fgColor='FFFF0000')
   clear_gray_block(s,6,2)
   self.assertEqual(s['D7'].value,'本文');self.assertIsNone(s['D7'].fill.patternType);self.assertEqual(s['F8'].fill.fgColor.rgb,'FFFF0000')
 def test_extended_slots_and_custom_rules_are_preserved_and_idempotent(self):
  s=openpyxl.Workbook().active
  for col in [2,12,82,202]:s.cell(6,col,'クラス')
  custom=DataValidation(type='list',formula1='"担当A,担当B"');custom.add('B14:C15');s.add_data_validation(custom)
  self.assertGreater(ensure_controls(s,self.reference()),0)
  for cell in ['CJ20','GZ20','CJ40','B14']:
   rules=[d for d in s.data_validations.dataValidation if cell in d.sqref];self.assertEqual(len(rules),1,cell)
  self.assertEqual([d.formula1 for d in s.data_validations.dataValidation if 'B14' in d.sqref],['"担当A,担当B"'])
  self.assertEqual(ensure_controls(s,self.reference()),0)
 def test_next_month_preserves_controls(self):
  w=openpyxl.Workbook();s=w.active;s.title='__TEMPLATE_MAIN__';s['B6']='クラス';s['GT6']='クラス'
  ensure_controls(s,self.reference());new=create_month_sheet(w,s,2026,10)
  assert_template_controls(new)
  self.assertTrue(any('GZ20' in d.sqref for d in new.data_validations.dataValidation))
 def test_repository_templates_are_valid(self):
  for name in ['TEMPLATE_MAIN.xlsx','TEMPLATE_X.xlsx']:
   w=openpyxl.load_workbook(Path(__file__).with_name(name));self.addCleanup(w.close);assert_template_controls(w.worksheets[0])
if __name__=='__main__':unittest.main()
